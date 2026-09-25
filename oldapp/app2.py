import streamlit as st
import sqlite3
import re
from docx import Document
from datetime import datetime
from zhipuai import ZhipuAI
import base64
import json
import io
from PIL import Image
import zipfile
import time
# ================= 配置区 =================
client = ZhipuAI(api_key="ac67cb8c61804fe6bdcecdc13e8aee0a.v7IK0Kxs4EMLsMbr")

# ================= 2. 工具函数 (兼顾数据库初始化) =================
def get_db_conn():
    conn = sqlite3.connect('school_review.db')
    conn.row_factory = sqlite3.Row
    try:
        conn.execute("ALTER TABLE submissions ADD COLUMN file_data BLOB")
    except sqlite3.OperationalError:
        pass
    # 新增：创建批次管理表
    conn.execute('''
        CREATE TABLE IF NOT EXISTS batches (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            batch_name TEXT UNIQUE,
            is_active INTEGER DEFAULT 1,
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
        )
    ''')
    return conn

# ================= 1. 初始化系统记忆 =================
if 'admin_logged_in' not in st.session_state:
    st.session_state.admin_logged_in = False
if 'grade_config' not in st.session_state:
    st.session_state.grade_config = {"高一": 38, "高二": 38, "高三": 37}
if 'current_selected_class' not in st.session_state:
    st.session_state.current_selected_class = None

# 动态获取当前正在进行的批次
conn = get_db_conn()
active_batch = conn.execute("SELECT batch_name FROM batches WHERE is_active=1 ORDER BY id DESC LIMIT 1").fetchone()
st.session_state.current_batch = active_batch['batch_name'] if active_batch else None
conn.close()

def get_class_status_from_db(batch, activity_type, grade):
    if not batch: return {}
    conn = get_db_conn()
    rows = conn.execute(
        "SELECT class_name, status, attempts FROM submissions WHERE batch_name=? AND activity_type=? AND grade=?",
        (batch, activity_type, grade)
    ).fetchall()
    conn.close()
    return {r['class_name']: {"status": r['status'], "attempts": r['attempts']} for r in rows}

def extract_images_from_docx(docx_file):
    images = []
    file_bytes = docx_file.read()
    try:
        with zipfile.ZipFile(io.BytesIO(file_bytes), 'r') as docx_zip:
            for item in docx_zip.namelist():
                if item.startswith('word/media/') and not item.endswith('/'):
                    try:
                        img_data = docx_zip.read(item)
                        if len(img_data) > 0: images.append(img_data)
                    except zipfile.BadZipFile: continue
    except Exception as e:
        raise Exception(f"文档结构损坏: {e}")
    return images

def check_time_duration(time_str):
    pattern = r'(\d{1,2})\s*[:：]\s*(\d{1,2})\s*[^\d]+\s*(\d{1,2})\s*[:：]\s*(\d{1,2})'
    match = re.search(pattern, time_str)
    if not match: return False, "未识别到有效时间段格式(需类似 15:36~15:48)"
    
    h1, m1, h2, m2 = map(int, match.groups())
    start = h1 * 60 + m1
    end = h2 * 60 + m2
    
    # 修复：先老老实实计算差值，再做判断
    diff = end - start
    if diff < 0: 
        diff += 24 * 60
        
    if diff < 10: 
        return False, f"时长不足10分钟(实际{diff}分钟)"
        
    return True, "时间合格"

def review_images_with_ai(images):
    try:
        pil_images = [Image.open(io.BytesIO(img)).convert("RGB").resize((600, 600)) for img in images]
        grid_img = Image.new('RGB', (1200, 1200), color='white')
        positions = [(0, 0), (600, 0), (0, 600), (600, 600)]
        for i, img in enumerate(pil_images[:4]): grid_img.paste(img, positions[i])
            
        buffer = io.BytesIO()
        grid_img.save(buffer, format="JPEG", quality=85)
        base64_img = base64.b64encode(buffer.getvalue()).decode('utf-8')
        
        content_list = [
            {"type": "text", "text": """
            审核4张照片拼成的四宫格。
            1. 每一张必须有时间水印(xx年xx日xx点xx分)。
            2. 右下角的图4必须是全员合照，人脸数必须大于3。
            返回JSON: {"watermark_passed": true/false, "faces_passed": true/false, "reason": "理由"}
            """},
            {"type": "image_url", "image_url": {"url": base64_img}}
        ]
        response = client.chat.completions.create(
            model="glm-4v", messages=[{"role": "user", "content": content_list}], temperature=0.1, timeout=40
        )
        result_text = response.choices[0].message.content.replace("```json", "").replace("```", "").strip()
        result = json.loads(result_text)
        return result.get("watermark_passed", False) and result.get("faces_passed", False), result.get("reason", "未知")
    except Exception as e:
        return False, f"AI审查失败: {e}"
# ================= 3. 侧边栏导航 =================
with st.sidebar:
    st.title("系统导航")
    page_mode = st.radio("请选择访问入口：", ["🟢 学生提交端", "⚙️ 管理员后台"])
    st.divider()
    st.caption(f"当前活动批次：{st.session_state.current_batch}")

# ================= 4. ⚙️ 管理员后台逻辑 =================
if page_mode == "⚙️ 管理员后台":
    st.title("⚙️ 团队活动中央控制台")
    if not st.session_state.admin_logged_in:
        pwd = st.text_input("管理员密码", type="password")
        if st.button("登录"):
            if pwd == "qiushihao2009":
                st.session_state.admin_logged_in = True
                st.rerun()
            else: st.error("密码错误！")
    else:
        if st.sidebar.button("登出系统"):
            st.session_state.admin_logged_in = False
            st.rerun()
            
        t1, t2, t3, t4 = st.tabs(["📊 综合统计", "📸 团日审核", "📦 导出", "⚙️ 设置"])
        with t1: 
            st.subheader(f"📊 {st.session_state.current_batch} - 综合统计大屏")
            stat_activity = st.radio("请选择要查看的活动类型：", ["🚩 团日活动", "🎈 志愿服务活动"], horizontal=True)
            current_stat_type = "团日活动" if "团日活动" in stat_activity else "志愿服务活动"
            
            conn = get_db_conn()
            rows = conn.execute(
                "SELECT grade, class_name, status FROM submissions WHERE batch_name=? AND activity_type=?", 
                (st.session_state.current_batch, current_stat_type)
            ).fetchall()
            conn.close()
            
            db_status_map = {f"{r['grade']}_{r['class_name']}": r['status'] for r in rows}
            st.divider()
            
            for grade in ["高一", "高二", "高三"]:
                total_classes = st.session_state.grade_config[grade]
                passed, pending, failed, unsubmitted = [], [], [], []
                
                for i in range(1, total_classes + 1):
                    cls_name = f"{i}班"
                    status = db_status_map.get(f"{grade}_{cls_name}", "white") 
                    if status == "green": passed.append(cls_name)
                    elif status == "pending": pending.append(cls_name)
                    elif status == "red": failed.append(cls_name)
                    else: unsubmitted.append(cls_name)
                    
                st.markdown(f"#### 📍 {grade} (已交: {len(passed) + len(pending)} / {total_classes} 班)")
                c1, c2, c3, c4 = st.columns(4)
                with c1: st.success(f"✅ 合格 ({len(passed)})\n\n" + "、".join(passed))
                with c2: st.warning(f"⏳ 待审 ({len(pending)})\n\n" + "、".join(pending))
                with c3: st.error(f"❌ 被驳回 ({len(failed)})\n\n" + "、".join(failed))
                with c4: 
                    st.info(f"⬜ 未交 ({len(unsubmitted)})")
                    with st.expander("名单"): st.write("、".join(unsubmitted))
                st.divider()

        with t2: 
            st.subheader("📸 团日活动 - 人工视觉审核池")
            conn = get_db_conn()
            pending_rows = conn.execute(
                "SELECT * FROM submissions WHERE activity_type='团日活动' AND status='pending' ORDER BY created_at ASC"
            ).fetchall()
            
            if len(pending_rows) == 0:
                st.success("🎉 太棒了！当前没有任何待审核的团日活动。")
            else:
                st.info(f"当前共有 {len(pending_rows)} 个班级等待人工审核。")
                if st.button("⚡ 一键通过所有待审班级 (特权操作)", type="primary"):
                    conn.execute("UPDATE submissions SET status='green' WHERE activity_type='团日活动' AND status='pending'")
                    conn.commit()
                    st.success("已清空所有待审任务，状态全部变更为合格！")
                    st.rerun()
                
                st.divider()
                current_task = pending_rows[0]
                task_id = current_task['id']
                st.write(f"正在审核：**{current_task['batch_name']} - {current_task['grade']} {current_task['class_name']}**")
                
                images = [base64.b64decode(img) for img in json.loads(current_task['images_data'])]
                c1, c2, c3 = st.columns(3)
                if len(images) == 3:
                    c1.image(images[0], use_container_width=True)
                    c2.image(images[1], use_container_width=True)
                    c3.image(images[2], use_container_width=True)
                
                col_btn1, col_btn2 = st.columns([1, 1])
                if col_btn1.button("✅ 画面合规 (通过)", key="pass_btn", use_container_width=True):
                    conn.execute("UPDATE submissions SET status='green' WHERE id=?", (task_id,))
                    conn.commit()
                    st.rerun()
                if col_btn2.button("❌ 画面违规 (驳回)", key="reject_btn", use_container_width=True):
                    conn.execute("UPDATE submissions SET status='red' WHERE id=?", (task_id,))
                    conn.commit()
                    st.rerun()
            conn.close()

        with t3: 
            st.subheader("📦 精准归档与分类导出")
            col1, col2 = st.columns(2)
            with col1: export_activity = st.selectbox("📌 选择活动类型", ["🚩 团日活动", "🎈 志愿服务活动"])
            with col2: export_grade = st.selectbox("🎓 选择年级", ["高一", "高二", "高三"])
                
            clean_activity = "团日活动" if "团日" in export_activity else "志愿服务活动"
            st.divider()
            
            if st.button(f"⚙️ 立即打包：【{export_grade} - {clean_activity}】", use_container_width=True, type="primary"):
                conn = get_db_conn()
                rows = conn.execute(
                    "SELECT class_name, file_data FROM submissions WHERE batch_name=? AND status='green' AND activity_type=? AND grade=?", 
                    (st.session_state.current_batch, clean_activity, export_grade)
                ).fetchall()
                conn.close()
                
                if len(rows) == 0:
                    st.warning(f"⚠️ 当前批次没有找到合格文件。")
                else:
                    with st.spinner("正在高速打包中..."):
                        zip_buffer = io.BytesIO()
                        with zipfile.ZipFile(zip_buffer, "a", zipfile.ZIP_DEFLATED, False) as zip_file:
                            for r in rows:
                                if r['file_data']:  
                                    zip_file.writestr(f"{r['class_name']}.docx", r['file_data'])
                        
                        st.success(f"🎉 打包成功！共提取了 {len(rows)} 份文件。")
                        dl_name = f"{st.session_state.current_batch}_{export_grade}_{clean_activity}.zip"
                        st.download_button("⬇️ 点击下载压缩包", data=zip_buffer.getvalue(), file_name=dl_name, mime="application/zip", use_container_width=True)

        with t4: 
            st.subheader("⚙️ 系统发布与历史归档")
            st.markdown("### 📢 当前收集项目")
            
            if st.session_state.current_batch:
                st.success(f"🟢 当前正在收集中：**{st.session_state.current_batch}**")
                
                if st.button("🛑 终止当前项目 (进入归档)", type="primary"):
                    conn = get_db_conn()
                    conn.execute("UPDATE batches SET is_active=0 WHERE batch_name=?", (st.session_state.current_batch,))
                    conn.commit()
                    conn.close()
                    st.session_state.current_batch = None
                    # 【核心魔法】：停顿0.2秒，让前端安全卸载，彻底避开报错
                    time.sleep(0.2) 
                    st.rerun()
            else:
                st.warning("⚪ 当前没有正在收集的项目。")
                with st.form("new_batch_form"):
                    new_batch = st.text_input("新建收集项目名称 (例如：2026年10月团日与志愿收集)：")
                    if st.form_submit_button("🚀 立即发布新项目"):
                        if new_batch:
                            try:
                                conn = get_db_conn()
                                conn.execute("INSERT INTO batches (batch_name, is_active) VALUES (?, 1)", (new_batch,))
                                conn.commit()
                                conn.close()
                                st.session_state.current_batch = new_batch
                                # 发布时同样停顿0.2秒
                                time.sleep(0.2)
                                st.rerun()
                            except sqlite3.IntegrityError:
                                st.error("该名称已存在，请更换一个项目名！")

            st.divider()
            st.markdown("### 🗄️ 历史项目记录")
            conn = get_db_conn()
            history = conn.execute("SELECT batch_name, created_at FROM batches WHERE is_active=0 ORDER BY id DESC").fetchall()
            
            if history:
                for h in history:
                    old_batch = h['batch_name']
                    with st.expander(f"📁 {old_batch} (发布于: {h['created_at']})"):
                        hist_activity = st.radio(f"查看 {old_batch} 的：", ["🚩 团日活动", "🎈 志愿服务活动"], key=f"radio_{old_batch}", horizontal=True)
                        clean_hist_act = "团日活动" if "团日" in hist_activity else "志愿服务活动"
                        
                        rows = conn.execute(
                            "SELECT grade, class_name, status FROM submissions WHERE batch_name=? AND activity_type=?", 
                            (old_batch, clean_hist_act)
                        ).fetchall()
                        
                        hist_status_map = {f"{r['grade']}_{r['class_name']}": r['status'] for r in rows}
                        
                        for grade in ["高一", "高二", "高三"]:
                            total_classes = st.session_state.grade_config[grade]
                            passed, pending, failed, unsubmitted = [], [], [], []
                            
                            for i in range(1, total_classes + 1):
                                cls_name = f"{i}班"
                                status = hist_status_map.get(f"{grade}_{cls_name}", "white") 
                                if status == "green": passed.append(cls_name)
                                elif status == "pending": pending.append(cls_name)
                                elif status == "red": failed.append(cls_name)
                                else: unsubmitted.append(cls_name)
                                
                            st.markdown(f"**📍 {grade}** (已交: {len(passed) + len(pending)} / {total_classes} 班)")
                            c1, c2, c3, c4 = st.columns(4)
                            with c1: st.success(f"✅ 合格 ({len(passed)})")
                            with c2: st.warning(f"⏳ 待审 ({len(pending)})")
                            with c3: st.error(f"❌ 驳回 ({len(failed)})")
                            with c4: st.info(f"⬜ 未交 ({len(unsubmitted)})")
                            st.divider()
            else:
                st.caption("暂无历史记录。")
            conn.close()
# ================= 5. 🟢 学生提交端逻辑 =================
# ================= 5. 🟢 学生提交端逻辑 =================
elif page_mode == "🟢 学生提交端":
    st.markdown("""
    <style>
    div.stButton > button {
        border-radius: 16px !important;
        border: 2px solid #f0f2f6 !important;
        font-weight: bold !important;
        transition: all 0.3s ease !important;
        padding: 10px !important;
    }
    div.stButton > button:hover {
        border-color: #ff6b81 !important;
        color: #ff6b81 !important;
        transform: translateY(-2px) !important;
        box-shadow: 0 6px 12px rgba(255,107,129,0.15) !important;
    }
    </style>
    """, unsafe_allow_html=True)

    st.title("✨ 班级活动自动化审查系统")
    
    if not st.session_state.current_batch:
        st.info("📭 当前暂无正在收集的活动项目，请等待管理员发布。")
    else:
        activity_type = st.radio("📌 请选择活动类型：", ["🎈 志愿服务活动", "🚩 团日活动"], horizontal=True)
        clean_act_type = "志愿服务活动" if "志愿" in activity_type else "团日活动"
        st.divider()
        
        grade_choice = st.radio("🎓 选择年级：", ["高一", "高二", "高三"], horizontal=True)
        st.subheader(f"📍 请选择 {grade_choice} 的班级")
        
        cols = st.columns(6)
        color_map = {"white": "⬜", "green": "✅", "red": "❌", "pending": "⏳"}
        current_db_status = get_class_status_from_db(st.session_state.current_batch, clean_act_type, grade_choice)
        
        for i in range(1, st.session_state.grade_config[grade_choice] + 1):
            cls_name = f"{i}班"
            data = current_db_status.get(cls_name, {"status": "white", "attempts": 0})
            btn_label = f"{color_map.get(data['status'], '⬜')} {cls_name}\n({data['attempts']}/3)"
            if cols[(i-1) % 6].button(btn_label, key=f"btn_{clean_act_type}_{grade_choice}_{cls_name}", use_container_width=True):
                st.session_state.active_selection = {"act": clean_act_type, "grade": grade_choice, "class": cls_name}
                
        current_sel = st.session_state.get('active_selection', {})
        if current_sel.get('act') == clean_act_type and current_sel.get('grade') == grade_choice:
            current_class = current_sel['class']
            class_info = current_db_status.get(current_class, {"status": "white", "attempts": 0})
            
            st.divider()
            st.subheader(f"📤 当前操作：{grade_choice} {current_class} - {clean_act_type}")
            
            if class_info["status"] == "green": 
                st.success("✅ 该班级已通过审查！")
            elif class_info["status"] == "pending": 
                st.warning("⏳ 已提交，正在等待管理员后台人工审核图片...")
            elif class_info["attempts"] >= 3: 
                st.error("🚫 提交次数已耗尽。")
            else:
                uploaded_file = st.file_uploader(f"上传《{clean_act_type}模板.docx》", type="docx")
                if uploaded_file is not None:
                    errors = []
                    try:
                        doc = Document(uploaded_file)
                        all_text = "".join([p.text for p in doc.paragraphs] + [cell.text for t in doc.tables for r in t.rows for cell in r.cells])
                        full_clean_text = re.sub(r'\s+', '', all_text).replace("：", ":").replace("、", "")
                        
                        if clean_act_type == "志愿服务活动":
                            participants_count = 0
                            time_duration_str = ""
                            nums = re.findall(r'参与人数.*?(\d+)', full_clean_text)
                            if nums: participants_count = int(nums[0])
                            
                            if "活动时长" in full_clean_text:
                                time_match = re.search(r'活动时长.{0,15}', full_clean_text)
                                if time_match: time_duration_str = time_match.group()
                                    
                            if participants_count == 0: errors.append("第2项错误：未读取到参与人数。")
                            elif grade_choice == "高一" and not (10 <= participants_count <= 12): errors.append(f"高一人数应为10-12人，实际 {participants_count} 人。")
                            elif grade_choice in ["高二", "高三"] and participants_count < 10: errors.append(f"{grade_choice}人数应≥10人，实际 {participants_count} 人。")
                            
                            phone_numbers = re.findall(r'1[3-9]\d{9}', full_clean_text)
                            if len(phone_numbers) != participants_count: errors.append(f"第4项错误：检测到的手机号数量({len(phone_numbers)}个)与填写的参与人数({participants_count}人)不一致。")
                                    
                            if not time_duration_str: errors.append("第3项错误：未读取到活动时长。")
                            else:
                                time_passed, time_msg = check_time_duration(time_duration_str)
                                if not time_passed: errors.append(f"第3项错误：{time_msg}")
                                
                            uploaded_file.seek(0)
                            images = extract_images_from_docx(uploaded_file)
                            if len(images) != 4: errors.append(f"第9项错误：图片必须为4张，实际提取到 {len(images)} 张。")
                            else:
                                st.info("✅ 正在呼叫 AI 审查图片，请稍候...")
                                ai_passed, ai_reason = review_images_with_ai(images)
                                if not ai_passed: errors.append(f"AI 图片审查未通过：{ai_reason}")
                                
                        elif clean_act_type == "团日活动":
                            uploaded_file.seek(0)
                            images = extract_images_from_docx(uploaded_file)
                            if len(images) != 3: errors.append(f"第九项错误：必须提交 3 张图片，实际 {len(images)} 张。")
                                
                            required_keywords = ["活动背景", "活动主题", "活动人数", "活动目的", "活动时间", "活动形式", "活动流程", "学生感想"]
                            for i, keyword in enumerate(required_keywords):
                                if keyword not in full_clean_text:
                                    errors.append(f"未找到【{keyword}】这一项标题。")
                                else:
                                    content_after = full_clean_text.split(keyword)[-1]
                                    content_between = content_after.split(required_keywords[i+1])[0] if i < len(required_keywords) - 1 else content_after
                                    pure_content = re.sub(r'[^\w\u4e00-\u9fa5]', '', content_between)
                                    if len(pure_content) < 2: errors.append(f"必填项【{keyword}】未填写有效内容。")

                        new_attempts = class_info["attempts"] + 1
                        uploaded_file.seek(0)
                        file_bytes = uploaded_file.read()
                        conn = get_db_conn()
                        conn.execute("DELETE FROM submissions WHERE batch_name=? AND activity_type=? AND grade=? AND class_name=?", (st.session_state.current_batch, clean_act_type, grade_choice, current_class))
                        
                        if len(errors) == 0:
                            final_status = "green" if clean_act_type == "志愿服务活动" else "pending"
                            images_b64 = [base64.b64encode(img).decode('utf-8') for img in images] if clean_act_type == "团日活动" else []
                            conn.execute('''INSERT INTO submissions (batch_name, activity_type, grade, class_name, status, attempts, images_data, file_data) VALUES (?, ?, ?, ?, ?, ?, ?, ?)''', (st.session_state.current_batch, clean_act_type, grade_choice, current_class, final_status, new_attempts, json.dumps(images_b64), file_bytes))
                            conn.commit()
                            st.success("🎉 审查/提交成功！")
                        else:
                            conn.execute('''INSERT INTO submissions (batch_name, activity_type, grade, class_name, status, attempts) VALUES (?, ?, ?, ?, ?, ?)''', (st.session_state.current_batch, clean_act_type, grade_choice, current_class, "red", new_attempts))
                            conn.commit()
                            st.error("❌ 审查未通过：\n" + "\n".join([f"{i+1}. {err}" for i, err in enumerate(errors)]))
                        conn.close()
                        if st.button("刷新看板状态"): st.rerun()
                    except Exception as e:
                        st.error(f"解析出错，文档可能损坏: {e}")