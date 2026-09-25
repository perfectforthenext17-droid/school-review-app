import os
os.environ["HTTP_PROXY"] = ""
os.environ["HTTPS_PROXY"] = ""

import streamlit as st
import re
from docx import Document
from datetime import datetime
from zhipuai import ZhipuAI
import base64
import json
import io
from PIL import Image
import zipfile
import time
from supabase import create_client
import gc

# ================= 全局屏蔽 ZIP CRC-32 严格校验 =================
_original_init = zipfile.ZipExtFile.__init__
def _patched_init(self, *args, **kwargs):
    _original_init(self, *args, **kwargs)
    self._expected_crc = None
zipfile.ZipExtFile.__init__ = _patched_init

# ================= 配置区 =================
client = ZhipuAI(api_key=st.secrets["ZHIPU_API_KEY"])

SUPABASE_URL = st.secrets["SUPABASE_URL"]
SUPABASE_KEY = st.secrets["SUPABASE_KEY"]
sb = create_client(SUPABASE_URL, SUPABASE_KEY)

# ================= 1. 初始化系统记忆 =================
if 'admin_logged_in' not in st.session_state:
    st.session_state.admin_logged_in = False
if 'grade_config' not in st.session_state:
    st.session_state.grade_config = {"高一": 38, "高二": 38, "高三": 37}
if 'current_selected_class' not in st.session_state:
    st.session_state.current_selected_class = None
if 'has_seen_announcement' not in st.session_state:
    st.session_state.has_seen_announcement = False
if 'announcement_data' not in st.session_state:
    st.session_state.announcement_data = None
if 'user_grade' not in st.session_state:
    st.session_state.user_grade = None
# ================= 2. 工具函数与状态查询 =================
@st.cache_data(ttl=60, show_spinner=False)
def get_active_batch():
    batch_res = sb.table("batches").select("batch_name").eq("is_active", 1).order("id", desc=True).limit(1).execute()
    return batch_res.data[0]['batch_name'] if batch_res.data else None

st.session_state.current_batch = get_active_batch()

@st.cache_data(ttl=15, show_spinner=False)
def get_class_status_from_db(batch, activity_type, grade):
    if not batch: return {}
    res = sb.table("submissions").select("class_name, status, attempts, failed_reason, file_path, original_filename").match({
        "batch_name": batch,
        "activity_type": activity_type,
        "grade": grade
    }).execute()
    return {r['class_name']: {
        "status": r['status'], 
        "attempts": r['attempts'],
        "failed_reason": r.get('failed_reason', ''),
        "file_path": r.get('file_path', ''),
        "original_filename": r.get('original_filename', '')
    } for r in res.data}

def extract_images_from_docx(docx_file):
    images = []
    file_bytes = docx_file.read()
    try:
        with zipfile.ZipFile(io.BytesIO(file_bytes), 'r') as docx_zip:
            for item in docx_zip.namelist():
                if item.startswith('word/media/') and item.lower().endswith(('.png', '.jpg', '.jpeg', '.bmp', '.webp')):
                    try:
                        img_data = docx_zip.read(item)
                        if len(img_data) > 0: images.append(img_data)
                    except Exception: 
                        continue
    except Exception as e:
        raise Exception(f"文档结构损坏: {e}")
    return images

def check_time_duration(time_str):
    pattern = r'(\d{1,2})\s*[:：]\s*(\d{1,2})\s*[^\d]+\s*(\d{1,2})\s*[:：]\s*(\d{1,2})'
    match = re.search(pattern, time_str)
    if not match: return False, "未识别到有效时间段格式(需类似 15:36~15:48)"
    h1, m1, h2, m2 = map(int, match.groups())
    start = h1 * 60 + m1
    end = h2 * 60 + m2
    diff = end - start
    if diff < 0: diff += 24 * 60
    if diff < 10: return False, f"时长不足10分钟(实际{diff}分钟)"
    return True, "时间合格"

def review_images_with_ai(images):
    try:
        pil_images = []
        for img_bytes in images:
            with Image.open(io.BytesIO(img_bytes)) as img:
                pil_images.append(img.convert("RGB").resize((600, 600)))
        
        grid_img = Image.new('RGB', (1200, 1200), color='white')
        positions = [(0, 0), (600, 0), (0, 600), (600, 600)]
        
        for i, img in enumerate(pil_images[:4]): 
            grid_img.paste(img, positions[i])
            img.close()
            
        buffer = io.BytesIO()
        grid_img.save(buffer, format="JPEG", quality=85)
        base64_img = base64.b64encode(buffer.getvalue()).decode('utf-8')
        
        grid_img.close()
        buffer.close()
        del pil_images
        del grid_img
        gc.collect() 
        
        content_list = [
            {"type": "text", "text": """
            审核4张照片拼成的四宫格。
            1. 每一张必须有时间水印(xx年xx日xx点xx分)。
            2. 右下角的图4必须是全员合照，人脸数必须大于3。
            返回JSON: {"watermark_passed": true/false, "faces_passed": true/false, "reason": "理由"}
            """},
            {"type": "image_url", "image_url": {"url": base64_img}}
        ]
        
        response = client.chat.completions.create(
            model="glm-4v", messages=[{"role": "user", "content": content_list}], temperature=0.1, timeout=40
        )
        
        del base64_img
        gc.collect()
        
        result_text = response.choices[0].message.content.replace("```json", "").replace("```", "").strip()
        result = json.loads(result_text)
        return result.get("watermark_passed", False) and result.get("faces_passed", False), result.get("reason", "未知")
    except Exception as e:
        return False, f"AI审查失败: {e}"

# ================= 3. 侧边栏导航 =================
with st.sidebar:
    st.title("系统导航")
    page_mode = st.radio("请选择访问入口：", ["🟢 学生提交端", "⚙️ 管理员后台"])
    st.divider()
    st.caption(f"当前活动批次：{st.session_state.current_batch}")

# ================= 4. ⚙️ 管理员后台逻辑 =================
if page_mode == "⚙️ 管理员后台":
    st.title("⚙️ 团队活动中央控制台")
    if not st.session_state.admin_logged_in:
        pwd = st.text_input("管理员密码", type="password")
        if st.button("登录"):
            if pwd == "qiushihao2009":
                st.session_state.admin_logged_in = True
                st.rerun()
            else: st.error("密码错误！")
    else:
        if st.sidebar.button("登出系统"):
            st.session_state.admin_logged_in = False
            st.rerun()
            
        t1, t2, t3, t4 = st.tabs(["📊 综合统计", "📸 团日审核", "📦 导出", "⚙️ 设置"])
        
        with t1: 
            st.subheader(f"📊 {st.session_state.current_batch} - 综合统计大屏")
            
            @st.dialog("📋 驳回详情查阅")
            def show_backend_rejected_details(class_data):
                st.markdown(f"### {class_data['grade']} {class_data['class_name']} - 驳回记录")
                st.divider()
                
                reason = class_data.get('failed_reason')
                if reason:
                    st.warning(f"**历史驳回原因：**\n\n{reason}")
                else:
                    st.warning("暂无详细驳回记录。")
                
                attempts = class_data.get('attempts', 0)
                if attempts >= 2 and class_data.get('file_path'):
                    st.success("🔒 **该班级 2 次机会已用尽，系统已锁定其最终错误原件。**")
                    with st.spinner("正在从云端飞速拉取锁定文件..."):
                        try:
                            doc_bytes = sb.storage.from_("school-docs").download(class_data['file_path'])
                            
                            st.download_button(
                                label=f"📥 提取锁定原件: {class_data.get('original_filename', '锁定文件.docx')}",
                                data=doc_bytes,
                                file_name=class_data.get('original_filename', '锁定文件.docx'),
                                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                                use_container_width=True
                            )
                            
                            with st.expander("👀 在线直接预览原件内容 (免下载打开)", expanded=False):
                                try:
                                    doc_io = io.BytesIO(doc_bytes)
                                    doc = Document(doc_io)
                                    
                                    st.markdown("**📄 提取文档文字内容：**")
                                    full_text = ""
                                    for para in doc.paragraphs:
                                        if para.text.strip(): full_text += para.text.strip() + "\n"
                                    for table in doc.tables:
                                        for row in table.rows:
                                            row_text = " | ".join([cell.text.strip() for cell in row.cells if cell.text.strip()])
                                            if row_text: full_text += row_text + "\n"
                                            
                                    st.text_area(label="已过滤掉排版格式的纯文本", value=full_text, height=200, disabled=True)
                                    
                                    st.markdown("**🖼️ 提取文档内部图片：**")
                                    doc_io.seek(0)
                                    preview_images = extract_images_from_docx(doc_io)
                                    if preview_images:
                                        img_cols = st.columns(len(preview_images) if len(preview_images) < 4 else 4)
                                        for idx, img_bytes in enumerate(preview_images):
                                            img_cols[idx % 4].image(img_bytes, use_container_width=True, caption=f"提取图片 {idx+1}")
                                    else:
                                        st.info("未在文档中检测到有效图片。")
                                except Exception as parse_e:
                                    st.error(f"在线解析预览失败，文档结构可能损坏: {parse_e}")
                        except Exception as e:
                            st.error(f"文件拉取失败，该记录可能为无文件强制通过产生的特殊记录: {e}")
                elif attempts == 1:
                    st.info("💡 **系统提示：** 该班级目前仅失败 1 次。根据系统容错机制，只记录了失败原因，并未将本次错误文件写入云端锁定，请等待该班级的最终提交。")

            stat_activity = st.radio("请选择要查看的活动类型：", ["🚩 团日活动", "🎈 志愿服务活动"], horizontal=True)
            current_stat_type = "团日活动" if "团日活动" in stat_activity else "志愿服务活动"
            
            res = sb.table("submissions").select("grade, class_name, status, attempts, failed_reason, file_path, original_filename").match({
                "batch_name": st.session_state.current_batch,
                "activity_type": current_stat_type
            }).execute()
            rows = res.data
            db_status_map = {f"{r['grade']}_{r['class_name']}": r for r in rows}
            st.divider()
            
            for grade in ["高一", "高二", "高三"]:
                total_classes = st.session_state.grade_config[grade]
                passed, pending, failed, unsubmitted = [], [], [], []
                
                for i in range(1, total_classes + 1):
                    cls_name = f"{i}班"
                    r_data = db_status_map.get(f"{grade}_{cls_name}", {"status": "white"}) 
                    status = r_data["status"]
                    
                    if status == "green": passed.append(cls_name)
                    elif status == "pending": pending.append(cls_name)
                    elif status == "red": failed.append(cls_name)
                    else: unsubmitted.append(cls_name)
                    
                st.markdown(f"#### 📍 {grade} (已交: {len(passed) + len(pending)} / {total_classes} 班)")
                c1, c2, c3, c4 = st.columns(4)
                with c1: st.success(f"✅ 合格 ({len(passed)})\n\n" + "、".join(passed))
                with c2: st.warning(f"⏳ 待审 ({len(pending)})\n\n" + "、".join(pending))
                with c3: 
                    st.error(f"❌ 驳回 ({len(failed)})")
                    if failed:
                        for f_class in failed:
                            btn_key = f"backend_fail_{st.session_state.current_batch}_{current_stat_type}_{grade}_{f_class}"
                            if st.button(f"🔍 查阅 {f_class}", key=btn_key, use_container_width=True):
                                show_backend_rejected_details(db_status_map[f"{grade}_{f_class}"])
                with c4: 
                    st.info(f"⬜ 未交 ({len(unsubmitted)})")
                    with st.expander("名单"): st.write("、".join(unsubmitted))
                st.divider()

        with t2: 
            st.subheader("📸 团日活动 - 人工视觉审核池")
            pending_res = sb.table("submissions").select("*").match({
                "activity_type": "团日活动",
                "status": "pending"
            }).order("created_at").execute()
            pending_rows = pending_res.data
            
            if len(pending_rows) == 0:
                st.success("🎉 太棒了！当前没有任何待审核的团日活动。")
            else:
                st.info(f"当前共有 {len(pending_rows)} 个班级等待人工审核。")
                if st.button("⚡ 一键通过所有待审班级 (特权操作)", type="primary"):
                    sb.table("submissions").update({"status": "green"}).match({"activity_type": "团日活动", "status": "pending"}).execute()
                    st.success("已清空所有待审任务，状态全部变更为合格！")
                    st.rerun()
                
                st.divider()
                current_task = pending_rows[0]
                task_id = current_task['id']
                st.write(f"正在审核：**{current_task['batch_name']} - {current_task['grade']} {current_task['class_name']}**")
                
                images = [base64.b64decode(img) for img in current_task['images_data']]
                c1, c2, c3 = st.columns(3)
                if len(images) == 3:
                    c1.image(images[0], use_container_width=True)
                    c2.image(images[1], use_container_width=True)
                    c3.image(images[2], use_container_width=True)
                
                col_btn1, col_btn2 = st.columns([1, 1])
                if col_btn1.button("✅ 画面合规 (通过)", key="pass_btn", use_container_width=True):
                    sb.table("submissions").update({"status": "green"}).eq("id", task_id).execute()
                    get_class_status_from_db.clear()
                    st.rerun()
                if col_btn2.button("❌ 画面违规 (驳回)", key="reject_btn", use_container_width=True):
                    sb.table("submissions").update({"status": "red"}).eq("id", task_id).execute()
                    get_class_status_from_db.clear()
                    st.rerun()

        with t3: 
            st.subheader("📦 精准归档与分类导出")
            st.info("📱 **移动端管理员须知**：由于微信/QQ内置浏览器限制，若无反应，请点击右上角「···」选择“在系统浏览器打开”后重试。")
            
            res_batches = sb.table("batches").select("batch_name").order("id", desc=True).execute()
            batch_rows = res_batches.data
            
            if not batch_rows:
                st.warning("⚠️ 系统中尚未创建任何批次记录。")
            else:
                batch_list = [row['batch_name'] for row in batch_rows]
                col_batch, col_act, col_grade = st.columns(3)
                with col_batch: export_batch = st.selectbox("📅 选择批次", batch_list)
                with col_act: export_activity = st.selectbox("📌 选择活动类型", ["🚩 团日活动", "🎈 志愿服务活动"])
                with col_grade: export_grade = st.selectbox("🎓 选择年级", ["高一", "高二", "高三"])
                    
                clean_activity = "团日活动" if "团日" in export_activity else "志愿服务活动"
                st.divider()
                task_key = f"{export_batch}_{export_grade}_{clean_activity}"
                
                if st.button(f"⚙️ 立即打包：【{export_batch} - {export_grade} - {clean_activity}】", use_container_width=True, type="primary"):
                    res_files = sb.table("submissions").select("class_name, file_path, original_filename").match({
                        "batch_name": export_batch,
                        "status": "green",
                        "activity_type": clean_activity,
                        "grade": export_grade
                    }).execute()
                    rows = res_files.data
                    
                    if len(rows) == 0:
                        st.warning(f"⚠️ 在【{export_batch}】批次下，没有找到【{export_grade} - {clean_activity}】的合格文件。")
                        if 'zip_data' in st.session_state: del st.session_state.zip_data
                    else:
                        with st.spinner("正在从云端飞速拉取文件并打包..."):
                            zip_buffer = io.BytesIO()
                            failed_downloads = []
                            with zipfile.ZipFile(zip_buffer, "a", zipfile.ZIP_DEFLATED, False) as zip_file:
                                def get_class_num(row):
                                    num_str = row['class_name'].replace('班', '')
                                    return int(num_str) if num_str.isdigit() else 999
                                sorted_rows = sorted(rows, key=get_class_num)
                                for r in sorted_rows:
                                    if r.get('file_path'):  
                                        try:
                                            doc_bytes = sb.storage.from_("school-docs").download(r['file_path'])
                                            og_name = r.get('original_filename')
                                            if og_name: file_name = og_name if r['class_name'] in og_name else f"{r['class_name']}_{og_name}"
                                            else: file_name = f"{r['class_name']}.docx"
                                            zip_file.writestr(file_name, doc_bytes)
                                        except Exception:
                                            failed_downloads.append(r['class_name'])
                                            continue
                            
                            st.session_state.zip_data = zip_buffer.getvalue()
                            st.session_state.zip_name = f"{task_key}.zip"
                            st.session_state.zip_count = len(sorted_rows) - len(failed_downloads)
                            if failed_downloads: st.error(f"❌ 以下班级拉取失败：{', '.join(failed_downloads)}")
                
                if st.session_state.get('zip_data') and st.session_state.get('zip_name') == f"{task_key}.zip":
                    st.success(f"🎉 打包成功！共提取了 {st.session_state.zip_count} 份文件。")
                    st.download_button(label="⬇️ 点击下载压缩包", data=st.session_state.zip_data, file_name=st.session_state.zip_name, mime="application/zip", use_container_width=True)

        with t4: 
            st.subheader("⚙️ 系统发布与历史归档")
            st.markdown("### 📢 当前收集项目")
            if st.session_state.current_batch:
                st.success(f"🟢 当前正在收集中：**{st.session_state.current_batch}**")
                if st.button("🛑 终止当前项目 (进入归档)", type="primary"):
                    sb.table("batches").update({"is_active": 0}).eq("batch_name", st.session_state.current_batch).execute()
                    st.session_state.current_batch = None
                    time.sleep(0.2) 
                    st.rerun()
            else:
                st.warning("⚪ 当前没有正在收集的项目。")
                with st.form("new_batch_form"):
                    new_batch = st.text_input("新建收集项目名称 (例如：2026年10月团日收集)：")
                    if st.form_submit_button("🚀 立即发布新项目"):
                        if new_batch:
                            try:
                                sb.table("batches").insert({"batch_name": new_batch, "is_active": 1}).execute()
                                st.session_state.current_batch = new_batch
                                time.sleep(0.2)
                                st.rerun()
                            except Exception: st.error("该名称已存在，请更换一个项目名！")

            st.divider()
            st.markdown("### 🗄️ 历史项目记录")
            hist_res = sb.table("batches").select("batch_name, created_at").eq("is_active", 0).order("id", desc=True).execute()
            history = hist_res.data
            if history:
                for h in history:
                    old_batch = h['batch_name']
                    with st.expander(f"📁 {old_batch} (发布于: {h['created_at']})"):
                        hist_activity = st.radio(f"查看 {old_batch} 的：", ["🚩 团日活动", "🎈 志愿服务活动"], key=f"radio_{old_batch}", horizontal=True)
                        clean_hist_act = "团日活动" if "团日" in hist_activity else "志愿服务活动"
                        
                        res_hist = sb.table("submissions").select("grade, class_name, status").match({"batch_name": old_batch, "activity_type": clean_hist_act}).execute()
                        hist_status_map = {f"{r['grade']}_{r['class_name']}": r['status'] for r in res_hist.data}
                        
                        for grade in ["高一", "高二", "高三"]:
                            total_classes = st.session_state.grade_config[grade]
                            passed, pending, failed, unsubmitted = [], [], [], []
                            for i in range(1, total_classes + 1):
                                cls_name = f"{i}班"
                                status = hist_status_map.get(f"{grade}_{cls_name}", "white") 
                                if status == "green": passed.append(cls_name)
                                elif status == "pending": pending.append(cls_name)
                                elif status == "red": failed.append(cls_name)
                                else: unsubmitted.append(cls_name)
                            st.markdown(f"**📍 {grade}** (已交: {len(passed) + len(pending)} / {total_classes} 班)")
                            c1, c2, c3, c4 = st.columns(4)
                            with c1: st.success(f"✅ 合格 ({len(passed)})")
                            with c2: st.warning(f"⏳ 待审 ({len(pending)})")
                            with c3: st.error(f"❌ 驳回 ({len(failed)})")
                            with c4: st.info(f"⬜ 未交 ({len(unsubmitted)})")
                            st.divider()
            else:
                st.caption("暂无历史记录。")

            # 👇 🆕 系统公告管理控制台 👇
            st.divider()
            st.markdown("### 📢 系统全局公告管理")
            st.info("在此编辑的公告，将在同学们初次进入前端系统时以【弹窗】形式强提醒弹出。")
            current_ann_text = ""
            try:
                ann_bytes = sb.storage.from_("school-docs").download("announcement.json")
                ann_data = json.loads(ann_bytes)
                if ann_data.get("is_active"): current_ann_text = ann_data.get("content", "")
            except Exception: pass
                
            with st.form("announcement_form"):
                new_announcement = st.text_area("编辑最新公告正文（清空文本并发布，即可关闭全站公告）：", value=current_ann_text, height=150)
                if st.form_submit_button("🚀 发布 / 更新公告"):
                    new_data = {"content": new_announcement, "is_active": bool(new_announcement.strip())}
                    sb.storage.from_("school-docs").upload("announcement.json", json.dumps(new_data).encode('utf-8'), file_options={"upsert": "true"})
                    st.session_state.announcement_data = new_data 
                    st.success("🎉 公告已全网同步更新！")
                    time.sleep(1)
                    st.rerun()

# ================= 5. 🟢 学生提交端逻辑 =================
elif page_mode == "🟢 学生提交端":
    import os, base64
    
    logo_base64, badge_base64, badge_mime = "", "", "jpeg"
    if os.path.exists("dff2f6bd9341d59fef8359f9cf1556f7.jpg"):
        with open("dff2f6bd9341d59fef8359f9cf1556f7.jpg", "rb") as f: logo_base64 = base64.b64encode(f.read()).decode('utf-8')
            
    # 🆕 智能侦测部门徽章：允许你将其简单命名为 badge.png 或 badge.jpg
    for file_name in ["badge.png", "badge.jpg", "badge.jpeg", "5ed956ea2d5b1f5dbad45fc99f7edf9c.jpg", "5ed956ea2d5b1f5dbad45fc99f7edf9c.png"]:
        if os.path.exists(file_name):
            with open(file_name, "rb") as f:
                badge_base64 = base64.b64encode(f.read()).decode('utf-8')
                # 自动匹配正确的图片底层格式
                badge_mime = "png" if file_name.lower().endswith('.png') else "jpeg"
            break
            
    logo_html = f'<img src="data:image/jpeg;base64,{logo_base64}" class="school-logo">' if logo_base64 else ''

    @st.dialog(" ")
    def show_announcement_dialog(content, badge_b64, mime):
        # 🖼️ 终极防白屏：如果真的没找到图片，自动降级显示一个高级护盾 Emoji
        img_html = f'<img src="data:image/{mime};base64,{badge_b64}" class="badge-img">' if badge_b64 else '<div class="badge-img fallback-badge">🛡️</div>'
        
        st.markdown(f"""
        <style>
        @keyframes badgePop {{ 0% {{ opacity: 0; transform: scale(0.6) translateY(-20px); }} 70% {{ transform: scale(1.08) translateY(0); }} 100% {{ opacity: 1; transform: scale(1) translateY(0); }} }}
        .badge-img {{ 
            width: 120px; height: 120px; border-radius: 50%; display: block; margin: 0 auto; 
            box-shadow: 0 10px 28px rgba(139, 28, 49, 0.25); border: 3px solid #ffffff; 
            animation: badgePop 0.8s cubic-bezier(0.2, 0.8, 0.2, 1) forwards;
            object-fit: cover; /* 保证无论上传长图还是方图，都会完美裁切成正圆 */
        }}
        .fallback-badge {{ display: flex; align-items: center; justify-content: center; font-size: 55px; background-color: var(--secondary-background-color); }}
        .announce-title {{ text-align: center; font-weight: 900; color: #8B1C31; font-size: 22px; margin-top: 18px; line-height: 1.4; border-bottom: 2px solid rgba(139, 28, 49, 0.1); padding-bottom: 12px; letter-spacing: 1px; }}
        .announce-content {{ margin-top: 18px; font-size: 15px; line-height: 1.8; color: var(--text-color); text-align: justify; padding: 0 10px; white-space: pre-wrap; }}
        @media (prefers-color-scheme: dark) {{ .badge-img {{ border: 3px solid #2b2b2b; }} }}
        </style>
        {img_html}
        <div class="announce-title">青年志愿者联合会<br>班级管理部公告</div>
        <div class="announce-content">{content}</div>
        """, unsafe_allow_html=True)
        st.divider()
        if st.button("✅ 我知道了", use_container_width=True, type="primary"):
            st.session_state.has_seen_announcement = True
            st.rerun()

    if not st.session_state.get('has_seen_announcement', False):
        if st.session_state.announcement_data is None:
            try:
                ann_bytes = sb.storage.from_("school-docs").download("announcement.json")
                st.session_state.announcement_data = json.loads(ann_bytes)
            except Exception:
                st.session_state.announcement_data = {"is_active": False, "content": ""}
        
        if st.session_state.announcement_data.get("is_active"):
            show_announcement_dialog(st.session_state.announcement_data["content"], badge_base64, badge_mime)
        else:
            # 💡 修复：如果当前没有开启公告，自动将其标记为已读，防止死锁
            st.session_state.has_seen_announcement = True

    # 👇 🆕 新增：防呆确认年级弹窗 👇
    @st.dialog("🎓 确认：请选择您的年级", width="small")
    def show_grade_selection_dialog():
        st.markdown("""
        <div style='text-align: center; margin-bottom: 15px;'>
            <span style='color: #8B1C31; font-weight: bold; font-size: 15px;'>
            ⚠️ 为防止默认选项导致提交错班级<br>请务必先手动确认您的所属年级：
            </span>
        </div>
        """, unsafe_allow_html=True)
        
        c1, c2, c3 = st.columns(3)
        if c1.button("高一", use_container_width=True, type="primary"):
            st.session_state.user_grade = "高一"
            st.rerun()
        if c2.button("高二", use_container_width=True, type="primary"):
            st.session_state.user_grade = "高二"
            st.rerun()
        if c3.button("高三", use_container_width=True, type="primary"):
            st.session_state.user_grade = "高三"
            st.rerun()

    st.markdown(f"""
    <style>
    :root {{ --dj-red: #8B1C31; --dj-red-light: rgba(139, 28, 49, 0.08); --dj-red-shadow: rgba(139, 28, 49, 0.25); --dj-card-bg: var(--secondary-background-color); --dj-text: var(--text-color); }}
    @media (prefers-color-scheme: dark) {{ :root {{ --dj-red: #ff768b; --dj-red-light: rgba(255, 118, 139, 0.15); --dj-red-shadow: rgba(255, 118, 139, 0.25); }} }}
    @keyframes fadeIn {{ from {{ opacity: 0; transform: translateY(12px); }} to {{ opacity: 1; transform: translateY(0); }} }}
    .block-container {{ animation: fadeIn 0.6s cubic-bezier(0.16, 1, 0.3, 1); }}
    @keyframes elegantEntrance {{ from {{ opacity: 0; transform: scale(0.9) translateY(10px); }} to {{ opacity: 1; transform: scale(1) translateY(0); }} }}
    .elegant-header {{ display: flex; align-items: center; justify-content: center; gap: 24px; padding: 20px 0 35px 0; margin-bottom: 10px; animation: elegantEntrance 1s cubic-bezier(0.2, 0.8, 0.2, 1) forwards; }}
    .school-logo {{ width: 120px; height: 120px; border-radius: 50%; object-fit: cover; box-shadow: 0 8px 18px rgba(139, 28, 49, 0.2); border: 3px solid var(--dj-card-bg); }}
    .header-text-container {{ display: flex; flex-direction: column; justify-content: center; }}
    .sub-title-tag {{ font-size: 14px; color: var(--dj-red); background-color: var(--dj-red-light); padding: 6px 16px; border-radius: 20px; font-weight: 700; letter-spacing: 1.5px; margin-bottom: 8px; width: fit-content; }}
    .main-title {{ font-size: 38px; font-weight: 900; color: var(--dj-text); margin: 0; line-height: 1.2; }}
    
    div[role="radiogroup"] {{ display: inline-flex !important; flex-direction: row !important; flex-wrap: wrap; background-color: var(--dj-card-bg) !important; padding: 6px !important; border-radius: 16px !important; gap: 4px !important; }}
    div[role="radiogroup"] label input[type="radio"] {{ display: none !important; }}
    div[role="radiogroup"] label input[type="radio"] + * {{ display: none !important; }}
    div[role="radiogroup"] label {{ padding: 10px 24px !important; border-radius: 12px !important; margin: 0 !important; background-color: transparent !important; transition: all 0.3s cubic-bezier(0.25, 0.8, 0.25, 1) !important; cursor: pointer; display: flex !important; align-items: center !important; justify-content: center !important; flex: 1; }}
    div[role="radiogroup"] label p {{ font-size: 16px !important; font-weight: 600 !important; color: var(--dj-text); margin: 0 !important; opacity: 0.6; transition: all 0.3s ease; }}
    div[role="radiogroup"] label:has(input:checked) {{ background-color: var(--dj-red) !important; box-shadow: 0 4px 14px var(--dj-red-shadow) !important; transform: scale(1.02); }}
    div[role="radiogroup"] label:has(input:checked) p {{ color: #ffffff !important; font-weight: 800 !important; opacity: 1; }}
    
    /* 1. 彻底隐藏状态锚点，兼容新老 Streamlit DOM */
    div[data-testid="element-container"]:has(.btn-status),
    div[data-testid="stElementContainer"]:has(.btn-status),
    .element-container:has(.btn-status) {{
        display: none !important; height: 0px !important; margin: 0px !important; padding: 0px !important;
    }}

    /* 2. 班级按钮高级质感 */
    div.stButton > button {{
        border-radius: 16px !important; font-weight: 700 !important; 
        transition: all 0.3s cubic-bezier(0.25, 0.8, 0.25, 1) !important; padding: 12px !important; 
    }}
    div.stButton > button:active {{ transform: scale(0.95) !important; }}

    /* ⚪ 未提交 (兼容双版本) */
    div[data-testid="element-container"]:has(.btn-status-white) + div[data-testid="element-container"] div.stButton > button,
    div[data-testid="stElementContainer"]:has(.btn-status-white) + div[data-testid="stElementContainer"] div.stButton > button {{
        background-color: var(--dj-card-bg) !important; color: var(--dj-text) !important; border: none !important;
    }}
    div[data-testid="element-container"]:has(.btn-status-white) + div[data-testid="element-container"] div.stButton > button:hover,
    div[data-testid="stElementContainer"]:has(.btn-status-white) + div[data-testid="stElementContainer"] div.stButton > button:hover {{
        transform: translateY(-2px) !important; box-shadow: 0 4px 12px rgba(0,0,0,0.08) !important;
    }}

    /* 🟢 已通过 (纯正实心 iOS Green) */
    div[data-testid="element-container"]:has(.btn-status-green) + div[data-testid="element-container"] div.stButton > button,
    div[data-testid="stElementContainer"]:has(.btn-status-green) + div[data-testid="stElementContainer"] div.stButton > button {{
        background-color: #34C759 !important; color: #ffffff !important; border: none !important; box-shadow: 0 4px 12px rgba(52, 199, 89, 0.3) !important;
    }}

    /* 🔴 审查驳回 (纯正实心 iOS Red) */
    div[data-testid="element-container"]:has(.btn-status-red) + div[data-testid="element-container"] div.stButton > button,
    div[data-testid="stElementContainer"]:has(.btn-status-red) + div[data-testid="stElementContainer"] div.stButton > button {{
        background-color: #FF3B30 !important; color: #ffffff !important; border: none !important; box-shadow: 0 4px 12px rgba(255, 59, 48, 0.3) !important;
    }}

    /* 🟡 待审状态 (纯正实心 iOS Orange) */
    div[data-testid="element-container"]:has(.btn-status-pending) + div[data-testid="element-container"] div.stButton > button,
    div[data-testid="stElementContainer"]:has(.btn-status-pending) + div[data-testid="stElementContainer"] div.stButton > button {{
        background-color: #FF9500 !important; color: #ffffff !important; border: none !important; box-shadow: 0 4px 12px rgba(255, 149, 0, 0.3) !important;
    }}
    
    @media (max-width: 768px) {{
        .elegant-header {{ flex-direction: column; text-align: center; gap: 16px; padding: 10px 0 25px 0; }}
        .header-text-container {{ align-items: center; }}
        .school-logo {{ width: 105px; height: 105px; }} 
        .main-title {{ font-size: 28px; }}
        .block-container {{ padding-left: 0.8rem !important; padding-right: 0.8rem !important; }}
        div.stButton > button {{ padding: 14px 2px !important; font-size: 14px !important; }}
        [data-testid="column"] {{ min-width: 30% !important; flex: 1 1 30% !important; padding: 0 6px !important; margin-bottom: 8px !important; }}
        div[role="radiogroup"] label {{ padding: 10px 12px !important; }}
        div[role="radiogroup"] label p {{ font-size: 14px !important; }}
    }}
    </style>
    
    <div class="elegant-header">
        {logo_html}
        <div class="header-text-container">
            <div class="sub-title-tag" style="margin-bottom: 6px;">东江中学 · 青志联班级管理部</div>
            <div class="sub-title-tag">奉献 友爱 互助 进步</div>
            <h1 class="main-title">班级活动审查系统</h1>
        </div>
    </div>
    """, unsafe_allow_html=True)
    
    if not st.session_state.current_batch:
        st.info("📭 当前暂无正在收集的活动项目，请等待管理员发布。")
    else:
        # 🛡️ 强制防呆屏障 1：如果公告还在展示，强制阻断下方内容渲染
        if not st.session_state.get('has_seen_announcement', False):
            st.stop()
            
        # 🛡️ 强制防呆屏障 2：如果还没选过年级，强制弹窗并阻断渲染
        if st.session_state.get('user_grade') is None:
            show_grade_selection_dialog()
            st.stop() # 强制中断！不让任何班级按钮显示出来，直到选完年级
            
        activity_type = st.radio("📌 请选择活动类型", ["志愿服务活动", "团日活动"], horizontal=True)
        clean_act_type = "志愿服务活动" if "志愿" in activity_type else "团日活动"
        st.divider()
        
        # 🧠 自动读取刚才弹窗里选好的年级，做成外围界面的“默认选项”
        grade_options = ["高一", "高二", "高三"]
        default_index = grade_options.index(st.session_state.user_grade)
        
        grade_choice = st.radio("🎓 选择年级", grade_options, index=default_index, horizontal=True)
        
        # 允许用户在选错后，依然可以在主界面上手动修改年级
        if grade_choice != st.session_state.user_grade:
            st.session_state.user_grade = grade_choice
            st.rerun()
            
        st.markdown(f"### 📍 请选择 {grade_choice} 的班级", unsafe_allow_html=True)
        
        current_db_status = get_class_status_from_db(st.session_state.current_batch, clean_act_type, grade_choice)
        total_classes = st.session_state.grade_config[grade_choice]
        
        @st.dialog("📤 班级活动文件自动审查与提交通道")
        def show_submission_dialog(clean_act_type, grade_choice, current_class, class_info):
            st.markdown(f"#### 当前操作：<span style='color: #8B1C31; font-weight: 900;'>{grade_choice} {current_class}</span> - {clean_act_type}", unsafe_allow_html=True)
            
            st.markdown("""
            <style>
            @keyframes floatWarning { 0% { transform: translateY(0px); box-shadow: 0 2px 8px rgba(255, 59, 48, 0.15); } 50% { transform: translateY(-4px); box-shadow: 0 8px 16px rgba(255, 59, 48, 0.4); } 100% { transform: translateY(0px); box-shadow: 0 2px 8px rgba(255, 59, 48, 0.15); } }
            .floating-warning-box { color: #FF3B30; font-weight: 900; font-size: 16px; text-align: center; padding: 12px; margin-top: -10px; margin-bottom: 12px; border: 2px solid #FF3B30; border-radius: 12px; background-color: rgba(255, 59, 48, 0.08); animation: floatWarning 2s ease-in-out infinite; letter-spacing: 1px; }
            </style>
            <div class="floating-warning-box">！！！文件上传次数只有两次，请检查无误后上传！！！</div>
            """, unsafe_allow_html=True)
            st.divider()

            if st.session_state.get('admin_logged_in', False):
                st.info("🛠️ **管理员特权操作台**")
                col_admin1, col_admin2 = st.columns(2)
                if col_admin1.button("🔄 次数归零", use_container_width=True):
                    sb.table("submissions").update({"attempts": 0}).match({"batch_name": st.session_state.current_batch, "activity_type": clean_act_type, "grade": grade_choice, "class_name": current_class}).execute()
                    get_class_status_from_db.clear()
                    st.rerun()
                if col_admin2.button("✅ 强制通过", use_container_width=True):
                    res = sb.table("submissions").select("id").match({"batch_name": st.session_state.current_batch, "activity_type": clean_act_type, "grade": grade_choice, "class_name": current_class}).execute()
                    if len(res.data) > 0:
                        sb.table("submissions").update({"status": "green"}).match({"batch_name": st.session_state.current_batch, "activity_type": clean_act_type, "grade": grade_choice, "class_name": current_class}).execute()
                    else:
                        sb.table("submissions").insert({"batch_name": st.session_state.current_batch, "activity_type": clean_act_type, "grade": grade_choice, "class_name": current_class, "status": "green", "attempts": 0, "images_data": [], "file_path": "", "original_filename": f"{current_class}_特批免检.docx"}).execute()
                    get_class_status_from_db.clear()
                    st.rerun()
                
                if class_info.get("failed_reason"):
                    st.warning(f"⚠️ **该班级历史驳回记录：**\n\n{class_info['failed_reason']}")
                if class_info.get("file_path"):
                    st.success("📦 **云端文件已就绪 (已通过 或 最终锁定)**")
                    with st.spinner("正在拉取文件流..."):
                        try:
                            doc_bytes = sb.storage.from_("school-docs").download(class_info["file_path"])
                            st.download_button(label=f"📥 提取文件: {class_info.get('original_filename', '归档文档.docx')}", data=doc_bytes, file_name=class_info.get("original_filename", "调阅文档.docx"), mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document", use_container_width=True)
                            with st.expander("👀 在线直接预览原件内容", expanded=False):
                                try:
                                    doc_io = io.BytesIO(doc_bytes)
                                    doc = Document(doc_io)
                                    full_text = ""
                                    for para in doc.paragraphs:
                                        if para.text.strip(): full_text += para.text.strip() + "\n"
                                    for table in doc.tables:
                                        for row in table.rows:
                                            row_text = " | ".join([cell.text.strip() for cell in row.cells if cell.text.strip()])
                                            if row_text: full_text += row_text + "\n"
                                    st.text_area(label="提取的纯文字", value=full_text, height=150, disabled=True)
                                    
                                    doc_io.seek(0)
                                    preview_images = extract_images_from_docx(doc_io)
                                    if preview_images:
                                        img_cols = st.columns(len(preview_images) if len(preview_images) < 4 else 4)
                                        for idx, img_bytes in enumerate(preview_images): img_cols[idx % 4].image(img_bytes, use_container_width=True)
                                    else:
                                        st.caption("未检测到有效图片。")
                                except Exception as parse_e: st.error(f"预览失败: {parse_e}")
                        except Exception as e:
                            st.error("拉取失败，可能是强制通过产生的空记录。")
                st.divider()

            if class_info["status"] == "green" or class_info["attempts"] >= 2:
                if class_info["status"] == "green": st.success("✅ 该班级已通过审查！无需再次提交。")
                else:
                    st.error("🚫 提交次数已耗尽，通道永久关闭。")
                    if not st.session_state.get('admin_logged_in', False) and class_info.get("failed_reason"):
                        st.warning(f"**最终驳回原因记录：**\n\n{class_info['failed_reason']}")
                
                if st.session_state.get('admin_logged_in', False):
                    st.divider()
                    st.info("👁️ **管理员特权：当前记录文件原件调阅**")
                    if class_info.get("file_path"):
                        try:
                            doc_bytes = sb.storage.from_("school-docs").download(class_info["file_path"])
                            st.download_button(label=f"📥 提取原件: {class_info.get('original_filename', '存底文档.docx')}", data=doc_bytes, file_name=class_info.get("original_filename", "存底文档.docx"), mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document", use_container_width=True)
                            with st.expander("👀 在线直接预览原件内容", expanded=False):
                                try:
                                    doc_io = io.BytesIO(doc_bytes)
                                    doc = Document(doc_io)
                                    full_text = ""
                                    for para in doc.paragraphs:
                                        if para.text.strip(): full_text += para.text.strip() + "\n"
                                    for table in doc.tables:
                                        for row in table.rows:
                                            row_text = " | ".join([cell.text.strip() for cell in row.cells if cell.text.strip()])
                                            if row_text: full_text += row_text + "\n"
                                    st.text_area(label="纯文本内容", value=full_text, height=150, disabled=True)
                                    doc_io.seek(0)
                                    preview_images = extract_images_from_docx(doc_io)
                                    if preview_images:
                                        img_cols = st.columns(len(preview_images) if len(preview_images) < 4 else 4)
                                        for idx, img_bytes in enumerate(preview_images): img_cols[idx % 4].image(img_bytes, use_container_width=True)
                                    else: st.caption("未检测到有效图片。")
                                except Exception as parse_e: st.error(f"预览失败: {parse_e}")
                        except Exception as dl_e: st.error(f"云端文件拉取失败: {dl_e}")
                    else: st.warning("该记录没有绑定任何云端文件。")
                        
            elif class_info["status"] == "pending": 
                st.warning("⏳ 文件已成功提交，正在等待管理员后台人工审核图片...")
                st.info("🔒 审核期间暂时锁定上传通道。如被驳回，可再次提交。")
            else:
                st.info("📱 **手机端提交指引**：请先在 WPS 或微信中将填好的文档“另存为/保存到手机本地”，然后再点击下方按钮上传。")
                uploaded_file = st.file_uploader(f"上传《{clean_act_type}模板.docx》", type="docx")
                if uploaded_file is not None:
                    errors = []
                    try:
                        doc = Document(uploaded_file)
                        all_text = "".join([p.text for p in doc.paragraphs] + [cell.text for t in doc.tables for r in t.rows for cell in r.cells])
                        full_clean_text = re.sub(r'\s+', '', all_text).replace("：", ":").replace("、", "")
                        
                        if clean_act_type == "志愿服务活动":
                            participants_count = 0
                            time_duration_str = ""
                            nums = re.findall(r'参与人数.*?(\d+)', full_clean_text)
                            if nums: participants_count = int(nums[0])
                            
                            if "活动时长" in full_clean_text:
                                time_match = re.search(r'活动时长.{0,15}', full_clean_text)
                                if time_match: time_duration_str = time_match.group()
                                    
                            if participants_count == 0: errors.append("第2项错误：未读取到参与人数。")
                            elif grade_choice == "高一" and not (10 <= participants_count <= 12): errors.append(f"高一人数应为10-12人，实际 {participants_count} 人。")
                            elif grade_choice in ["高二", "高三"] and participants_count < 10: errors.append(f"{grade_choice}人数应≥10人，实际 {participants_count} 人。")
                            
                            phone_numbers = re.findall(r'1[3-9]\d{9}', full_clean_text)
                            if len(phone_numbers) != participants_count: errors.append(f"第4项错误：检测到的手机号数量({len(phone_numbers)}个)与填写的参与人数({participants_count}人)不一致。")
                                    
                            if not time_duration_str: errors.append("第3项错误：未读取到活动时长。")
                            else:
                                time_passed, time_msg = check_time_duration(time_duration_str)
                                if not time_passed: errors.append(f"第3项错误：{time_msg}")
                                
                            uploaded_file.seek(0)
                            images = extract_images_from_docx(uploaded_file)
                            if len(images) != 4: errors.append(f"第9项错误：图片必须为4张，实际提取到 {len(images)} 张。")
                            else:
                                with st.spinner("视觉引擎高速运转中... 正在进行多模态 AI 图像审查 (约5~10秒)"):
                                    ai_passed, ai_reason = review_images_with_ai(images)
                                if not ai_passed: errors.append(f"AI 图片审查未通过：{ai_reason}")
                                
                        elif clean_act_type == "团日活动":
                            uploaded_file.seek(0)
                            images = extract_images_from_docx(uploaded_file)
                            # 如果确实是10张图，这里依然会准确拦截。只有3张图才能通过。
                            if len(images) != 3: errors.append(f"第九项错误：必须提交 3 张图片，实际 {len(images)} 张。")
                                
                            required_keywords = ["活动背景", "活动主题", "活动人数", "活动目的", "活动时间", "活动形式", "活动流程", "学生感想"]
                            for i, keyword in enumerate(required_keywords):
                                # 💡 兼容性扩展：如果找不到“学生感想”，允许学生使用相近的表述
                                search_keyword = keyword
                                if keyword == "学生感想" and keyword not in full_clean_text:
                                    for alt_k in ["团员感想", "个人感想", "心得体会", "活动感想", "感想"]:
                                        if alt_k in full_clean_text:
                                            search_keyword = alt_k
                                            break
                                            
                                if search_keyword not in full_clean_text:
                                    errors.append(f"未找到【{keyword}】这一项标题。")
                                else:
                                    # 1. 找到该关键字【第一次】出现之后的所有文本，防止抓到末尾空壳
                                    content_after = full_clean_text.split(search_keyword, 1)[1]
                                    
                                    # 2. 寻找这段文本中，最先出现的【其他任意关键字】，将其作为精准截断点
                                    next_positions = []
                                    for other_k in required_keywords:
                                        if other_k != keyword and other_k in content_after:
                                            next_positions.append(content_after.find(other_k))
                                            
                                    if next_positions:
                                        content_between = content_after[:min(next_positions)]
                                    else:
                                        content_between = content_after
                                        
                                    pure_content = re.sub(r'[^\w\u4e00-\u9fa5]', '', content_between)
                                    if len(pure_content) < 2: errors.append(f"必填项【{keyword}】未填写有效内容。")

                        new_attempts = class_info["attempts"] + 1
                        uploaded_file.seek(0)
                        file_bytes = uploaded_file.read()
                        
                        sb.table("submissions").delete().match({"batch_name": st.session_state.current_batch, "activity_type": clean_act_type, "grade": grade_choice, "class_name": current_class}).execute()

                        if len(errors) == 0:
                            final_status = "green" if clean_act_type == "志愿服务活动" else "pending"
                            images_b64 = [base64.b64encode(img).decode('utf-8') for img in images] if clean_act_type == "团日活动" else []
                            
                            safe_storage_path = f"doc_{int(time.time() * 1000)}.docx"
                            sb.storage.from_("school-docs").upload(path=safe_storage_path, file=file_bytes, file_options={"upsert": "true"})
                            
                            sb.table("submissions").insert({
                                "batch_name": st.session_state.current_batch, "activity_type": clean_act_type, "grade": grade_choice, "class_name": current_class,
                                "status": final_status, "attempts": new_attempts, "images_data": images_b64, "file_path": safe_storage_path, "original_filename": uploaded_file.name
                            }).execute()
                            
                            st.success("🎉 审查/提交成功！即将自动返回...")
                            time.sleep(2)
                            st.rerun() 
                        else:
                            error_text = "\n".join([f"{i+1}. {err}" for i, err in enumerate(errors)])
                            if new_attempts == 1:
                                reason = f"【第1次尝试失败】\n{error_text}"
                                sb.table("submissions").insert({"batch_name": st.session_state.current_batch, "activity_type": clean_act_type, "grade": grade_choice, "class_name": current_class, "status": "red", "attempts": new_attempts, "failed_reason": reason}).execute()
                                st.error(f"❌ 审查未通过（这是您第1次提交）：\n{error_text}\n\n**系统提示：您还有最后 1 次更正机会，本次错误文件未被后台收录入库。**")
                            else:
                                old_reason = class_info.get("failed_reason") or ""
                                reason = old_reason + f"\n\n【第2次尝试失败 (绝路锁定)】\n{error_text}"
                                safe_storage_path = f"doc_locked_{int(time.time() * 1000)}.docx"
                                sb.storage.from_("school-docs").upload(path=safe_storage_path, file=file_bytes, file_options={"upsert": "true"})
                                sb.table("submissions").insert({"batch_name": st.session_state.current_batch, "activity_type": clean_act_type, "grade": grade_choice, "class_name": current_class, "status": "red", "attempts": new_attempts, "failed_reason": reason, "file_path": safe_storage_path, "original_filename": f"[二次失败强制存底]_{uploaded_file.name}"}).execute()
                                st.error(f"❌ 审查再次未通过：\n{error_text}\n\n**系统锁定提示：2次机会已用尽，该错误文件已被强制上传并作为最终凭证入库，当前班级通道永久关闭。**")
                            
                            if st.button("确认并重试/退出"): st.rerun()
                                
                    except Exception as e:
                        st.error(f"解析出错，文档可能损坏: {e}")
                    finally:
                        get_class_status_from_db.clear()
                        if 'file_bytes' in locals(): del file_bytes
                        if 'images' in locals(): del images
                        if 'images_b64' in locals(): del images_b64
                        if 'all_text' in locals(): del all_text
                        import gc
                        gc.collect() 
        
        for i in range(1, total_classes + 1, 6):
            cols = st.columns(6)
            for j in range(6):
                class_idx = i + j
                if class_idx <= total_classes:
                    cls_name = f"{class_idx}班"
                    data = current_db_status.get(cls_name, {"status": "white", "attempts": 0})
                    col = cols[j]
                    col.markdown(f"<div class='btn-status btn-status-{data['status']}'></div>", unsafe_allow_html=True)
                    btn_label = f"{cls_name}\n({data['attempts']}/2)"
                    if col.button(btn_label, key=f"btn_{clean_act_type}_{grade_choice}_{cls_name}", use_container_width=True):
                        show_submission_dialog(clean_act_type, grade_choice, cls_name, data)

# ================= 全局网页最底部声明 =================
st.divider()
st.markdown("""
<div style="text-align: center; color: var(--dj-text); opacity: 0.45; font-size: 12px; padding: 20px 0 40px 0;">
    本网站由 邱世豪 与 Google Gemini 协作开发 | 出现问题欢迎积极反馈
</div>
""", unsafe_allow_html=True)