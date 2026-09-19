from docx import Document

print("=== 开始提取【团日活动模板】===")
doc_tuanri = Document("团日活动模板.docx")
for para in doc_tuanri.paragraphs:
    # 只要段落里有字，就打印出来
    if para.text.strip():
        print(para.text)

print("\n=== 开始提取【志愿服务活动模板】表格 ===")
doc_zhiyuan = Document("志愿服务活动模板.docx")
if doc_zhiyuan.tables:
    table = doc_zhiyuan.tables[0] # 抓取第一个表格
    for i, row in enumerate(table.rows):
        # 把表格每一行的内容提取出来，用 | 隔开
        row_text = " | ".join([cell.text.strip() for cell in row.cells])
        print(f"第 {i+1} 行: {row_text}")
else:
    print("没有找到表格")
    