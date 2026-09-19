from docx import Document

def test_read_file(file_name):
    try:
        # 尝试打开文档
        doc = Document(file_name)
        print(f"✅ 成功读取文件: {file_name}")
        print(f"   该文档共有 {len(doc.paragraphs)} 个段落，以及 {len(doc.tables)} 个表格。")
        print("-" * 40)
    except Exception as e:
        print(f"❌ 读取 {file_name} 失败，请检查文件名是否正确。错误信息: {e}")

# 测试读取你的两个模板
test_read_file("团日活动模板.docx")
test_read_file("志愿服务活动模板.docx")